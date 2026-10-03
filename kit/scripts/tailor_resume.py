#!/usr/bin/env python3
"""Tailor the master resume for one posting and build the application files (job-agent-kit).

Requires python-docx, pypdf and pyyaml (installed into ~/job-search/.venv by install.sh).
Run from ~/job-search:

  .venv/bin/python scripts/tailor_resume.py --profile profile/profile.yaml \
      --spec applications/2026-10-02/li-4012345678/spec.json [--no-pdf] [--strict-numbers]
  .venv/bin/python scripts/tailor_resume.py --dump profile/master-resume.docx [--json]

spec.json
  {
    "company": "Example AI", "role": "Forward Deployed Engineer",
    "job_url": "https://...",
    "out_dir": "optional; relative paths resolve against the current directory; default = spec folder",
    "file_stem": "optional; default '<Name> - Resume - <Company>'",
    "replacements": [{"anchor": "headline", "text": "..."}, {"anchor": "summary_1", "text": "..."}],
    "cover_letter": {"greeting": "Dear Hiring Team,", "paragraphs": ["...", "..."], "closing": "Best regards,"}
  }

Each anchor name refers to resume.anchors.<name> in profile.yaml: the exact current text of a
paragraph in the master resume. Matching ignores differences in whitespace, quotes and dashes.
A paragraph whose whole text equals the anchor is replaced (new text goes into the first run,
keeping its formatting; the other runs are cleared). Otherwise a paragraph that contains the
anchor has just that substring replaced, provided the anchor is at least 12 characters and sits
at the start or end of the paragraph or covers at least half of it (the text after a bold
"Skills:" label is the typical case). Body, tables, text boxes, headers and footers are searched.

Exit codes: 0 ok; 1 usage, spec or profile error; 2 anchor not found or ambiguous (no files written);
3 resume PDF longer than resume.max_pages; 4 PDF conversion failed; 5 unsupported numbers with
--strict-numbers. A JSON summary is printed to stdout; messages go to stderr.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    import docx
    import yaml
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from pypdf import PdfReader
except ImportError as exc:  # pragma: no cover - environment problem
    sys.stderr.write(
        f"ERROR: missing Python package ({exc.name}). Run the scripts with the workspace Python:\n"
        "  .venv/bin/python scripts/tailor_resume.py ...\n"
        "or install the packages: .venv/bin/python -m pip install python-docx pypdf pyyaml\n"
    )
    sys.exit(1)

EXIT_OK, EXIT_USAGE, EXIT_ANCHOR, EXIT_PAGES, EXIT_PDF, EXIT_NUMBERS = 0, 1, 2, 3, 4, 5
HEADLINE_MAX = 85

W_P, W_T, W_TAB, W_BR, W_CR = qn("w:p"), qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")
W_NBH, W_PPR, W_TXBX, W_TC = qn("w:noBreakHyphen"), qn("w:pPr"), qn("w:txbxContent"), qn("w:tc")
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
SPECIAL_TEXT = {W_TAB: "\t", W_BR: "\n", W_CR: "\n", W_NBH: "-"}
REL_HEADER = "/relationships/header"
REL_FOOTER = "/relationships/footer"

QUOTES = {
    "\u2018": "'", "\u2019": "'", "\u201a": "'", "\u201b": "'", "\u2032": "'", "\u00b4": "'", "`": "'",
    "\u201c": '"', "\u201d": '"', "\u201e": '"', "\u201f": '"', "\u2033": '"', "\u00ab": '"', "\u00bb": '"',
}
DASHES = {c: "-" for c in "\u2010\u2011\u2012\u2013\u2014\u2015\u2212"}
INVISIBLE = set("\u200b\u200c\u200d\u2060\ufeff\u00ad")


def log(msg: str) -> None:
    sys.stderr.write(msg + "\n")


# ----------------------------------------------------------------- normalization

def normalize_with_map(text: str) -> tuple[str, list[int]]:
    """Normalize whitespace, quotes and dashes; return the text and, per output char, its source index."""
    out: list[str] = []
    src: list[int] = []
    prev_space = True  # drops leading whitespace
    for i, ch in enumerate(text):
        if ch in INVISIBLE:
            continue
        if ch.isspace():
            if not prev_space:
                out.append(" ")
                src.append(i)
            prev_space = True
            continue
        if ch == "\u2026":
            out.extend("...")
            src.extend([i, i, i])
        else:
            out.append(QUOTES.get(ch) or DASHES.get(ch) or ch)
            src.append(i)
        prev_space = False
    if out and out[-1] == " ":
        out.pop()
        src.pop()
    return "".join(out), src


def normalize(text: str) -> str:
    return normalize_with_map(text)[0]


# ----------------------------------------------------------------- paragraph walk

class Para:
    """A w:p element with its text segments (w:t, tabs, breaks) in document order."""

    def __init__(self, element, location: str):
        self.el = element
        self.location = location

    def segments(self) -> list[tuple[object, str]]:
        segs: list[tuple[object, str]] = []

        def walk(node) -> None:
            for child in node:
                tag = child.tag
                if tag in (W_PPR, W_TXBX, W_P):  # paragraph props; nested text-box paragraphs are walked separately
                    continue
                if tag == W_T:
                    segs.append((child, child.text or ""))
                elif tag in SPECIAL_TEXT:
                    segs.append((child, SPECIAL_TEXT[tag]))
                else:
                    walk(child)

        walk(self.el)
        return segs

    @property
    def text(self) -> str:
        return "".join(t for _, t in self.segments())

    @property
    def style(self) -> str:
        ppr = self.el.find(W_PPR)
        if ppr is not None:
            ps = ppr.find(qn("w:pStyle"))
            if ps is not None:
                return ps.get(qn("w:val")) or ""
        return ""

    @property
    def primary(self) -> bool:
        """False for the legacy (mc:Fallback) copy of a text box that Word stores twice."""
        return self.location != "textbox-fallback"


def locate(element, part_label: str) -> str:
    in_txbx = in_fallback = in_cell = False
    for anc in element.iterancestors():
        if anc.tag == W_TXBX:
            in_txbx = True
        elif anc.tag == MC_FALLBACK:
            in_fallback = True
        elif anc.tag == W_TC:
            in_cell = True
    if in_txbx:
        return "textbox-fallback" if in_fallback else "textbox"
    if in_cell:
        return "table" if part_label == "body" else f"{part_label}-table"
    return part_label


def all_paragraphs(document) -> list[Para]:
    roots = [("body", document.element.body)]
    seen = set()
    for rel in document.part.rels.values():
        if rel.is_external:
            continue
        label = "header" if rel.reltype.endswith(REL_HEADER) else "footer" if rel.reltype.endswith(REL_FOOTER) else None
        if label and id(rel.target_part) not in seen:
            seen.add(id(rel.target_part))
            roots.append((label, rel.target_part.element))
    return [Para(p, locate(p, label)) for label, root in roots for p in root.iter(W_P)]


# ----------------------------------------------------------------- replacement

def set_text(t_el, text: str) -> None:
    t_el.text = text
    t_el.set(XML_SPACE, "preserve")


def replace_range(para: Para, start: int, end: int, new_text: str) -> None:
    """Replace original-text range [start, end) of the paragraph with new_text.

    The run holding the first matched character receives the new text (so its formatting is
    kept); other matched text is cleared; tabs and breaks inside the range are removed.
    """
    pos, host_done = 0, False
    for el, text in para.segments():
        s, e = pos, pos + len(text)
        pos = e
        if not text or e <= start or s >= end:
            continue
        if el.tag == W_T:
            suffix = text[end - s:] if e > end else ""
            if not host_done:
                set_text(el, text[: max(0, start - s)] + new_text + suffix)
                host_done = True
            else:
                set_text(el, suffix)
        else:
            el.getparent().remove(el)
    if not host_done:
        raise RuntimeError("no text run inside the matched range")


MIN_SUBSTRING = 12  # shorter anchors may only match a whole paragraph


def find_matches(paras: list[Para], anchor: str) -> tuple[str, list[tuple[Para, int, int]]]:
    """Return (kind, [(para, orig_start, orig_end)]); kind is exact, substring, ambiguous, fragment or none.

    Exact (whole-paragraph) matches win. The substring fallback only accepts an anchor of at least
    MIN_SUBSTRING characters that sits at the start or end of its paragraph (the text after a
    "Skills:" label) or covers at least half of it, so a careless short anchor cannot rewrite a
    fragment in the middle of an unrelated bullet.
    """
    target = normalize(anchor)
    exact, partial, fragments = [], [], 0
    for para in paras:
        norm, src = normalize_with_map(para.text)
        if not norm:
            continue
        if norm == target:
            exact.append((para, src[0], src[-1] + 1))
            continue
        idx = norm.find(target)
        if idx < 0:
            continue
        at_edge = idx == 0 or idx + len(target) == len(norm)
        if len(target) < MIN_SUBSTRING or not (at_edge or len(target) * 2 >= len(norm)):
            fragments += para.primary
            continue
        if norm.find(target, idx + 1) >= 0:
            partial.append((para, -1, -1))  # repeated inside one paragraph: ambiguous
        else:
            partial.append((para, src[idx], src[idx + len(target) - 1] + 1))
    if exact:
        return "exact", exact
    primary = [m for m in partial if m[0].primary]
    if len(primary) == 1 and primary[0][1] >= 0:
        fallback_copies = [m for m in partial if not m[0].primary and m[1] >= 0]
        return "substring", primary + fallback_copies
    if partial:
        return "ambiguous", partial
    return ("fragment", []) if fragments else ("none", [])


# ----------------------------------------------------------------- numbers guard

NUMBER_RE = re.compile(
    r"(?<![\w.\-/])(?:\$|USD\s?|US\$|EUR\s?|\u20ac|\u00a3)?"
    r"(\d{1,3}(?:,\d{3})+|\d+)(\.\d+)?"
    r"(?:\s?(%|percent\b|x\b|\u00d7|k\b|K\b|mm\b|MM\b|m\b|M\b|bn\b|b\b|B\b|million\b|billion\b|thousand\b))?"
    r"\+?(?![\w])"
)
MULTIPLIER = {"k": 1000, "thousand": 1000, "m": 10**6, "mm": 10**6, "million": 10**6, "b": 10**9, "bn": 10**9, "billion": 10**9}


def number_tokens(text: str) -> list[tuple[str, str]]:
    """Return (as written, canonical) pairs, e.g. ('$1.2M', '1200000') or ('38 percent', '38%')."""
    out = []
    for m in NUMBER_RE.finditer(text or ""):
        unit = (m.group(3) or "").lower()
        try:
            value = Decimal(m.group(1).replace(",", "") + (m.group(2) or ""))
        except InvalidOperation:
            continue
        kind = ""
        if unit in MULTIPLIER:
            value *= MULTIPLIER[unit]
        elif unit in ("%", "percent"):
            kind = "%"
        elif unit in ("x", "\u00d7"):
            kind = "x"
        canon = format(value.normalize(), "f")
        if "." in canon:
            canon = canon.rstrip("0").rstrip(".")
        out.append((m.group(0).strip(), canon + kind))
    return out


def numbers_guard(spec: dict, source_text: str) -> dict:
    known = {c for _, c in number_tokens(source_text)}
    places = [(r["anchor"], r["text"]) for r in spec["replacements"]]
    cover = spec.get("cover_letter") or {}
    places += [(f"cover_letter.paragraphs[{i}]", p) for i, p in enumerate(cover.get("paragraphs") or [])]
    checked, unsupported = 0, []
    for where, text in places:
        for written, canon in number_tokens(text):
            checked += 1
            if canon not in known:
                unsupported.append({"where": where, "number": written})
    return {"checked": checked, "unsupported": unsupported}


# ----------------------------------------------------------------- inputs

def fail(code: int, msg: str, summary: dict | None = None):
    log(f"ERROR: {msg}")
    if summary is not None:
        summary.update(ok=False, error=msg, exit_code=code)
        print(json.dumps(summary, indent=2))
    sys.exit(code)


def load_yaml(path: Path) -> dict:
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        fail(EXIT_USAGE, f"profile not found: {path}")
    except yaml.YAMLError as exc:
        fail(EXIT_USAGE, f"profile.yaml is not valid YAML: {exc}")
    if not isinstance(data, dict):
        fail(EXIT_USAGE, f"{path} must be a YAML mapping")
    return data


def resolve_master(profile_path: Path, master: str) -> Path:
    raw = Path(os.path.expanduser(str(master)))
    candidates = [raw] if raw.is_absolute() else [
        Path.cwd() / raw, profile_path.resolve().parent.parent / raw, profile_path.resolve().parent / raw,
    ]
    for c in candidates:
        if c.is_file():
            return c
    fail(EXIT_USAGE, f"resume.master_path not found: {master} (looked in: {', '.join(str(c) for c in candidates)})")


def load_spec(path: Path) -> dict:
    try:
        spec = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        fail(EXIT_USAGE, f"spec not found: {path}")
    except json.JSONDecodeError as exc:
        fail(EXIT_USAGE, f"spec.json is not valid JSON: {exc}")
    if not isinstance(spec, dict):
        fail(EXIT_USAGE, "spec.json must be a JSON object")
    for key in ("company", "role"):
        if not str(spec.get(key) or "").strip():
            fail(EXIT_USAGE, f"spec.json needs a non-empty '{key}'")
    reps = spec.get("replacements")
    if not isinstance(reps, list) or not reps:
        fail(EXIT_USAGE, "spec.json needs a non-empty 'replacements' list")
    seen = set()
    for i, r in enumerate(reps):
        if not isinstance(r, dict) or not str(r.get("anchor") or "").strip() or not str(r.get("text") or "").strip():
            fail(EXIT_USAGE, f"replacements[{i}] needs non-empty 'anchor' and 'text'")
        if r["anchor"] in seen:
            fail(EXIT_USAGE, f"anchor '{r['anchor']}' appears twice in replacements")
        seen.add(r["anchor"])
        r["text"] = " ".join(str(r["text"]).split())
    cover = spec.get("cover_letter")
    if cover is not None:
        if not isinstance(cover, dict) or not isinstance(cover.get("paragraphs"), list) or not cover["paragraphs"]:
            fail(EXIT_USAGE, "cover_letter needs a non-empty 'paragraphs' list")
    return spec


def safe_name(text: str) -> str:
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", str(text))
    return re.sub(r"\s+", " ", text).strip().rstrip(".") or "untitled"


# ----------------------------------------------------------------- cover letter

def contact_line(identity: dict) -> str:
    bits = [identity.get(k) for k in ("email", "phone", "location")]
    links = identity.get("links") or {}
    links = list(links.values()) if isinstance(links, dict) else list(links)
    bits += [re.sub(r"^https?://(www\.)?", "", str(link).strip()).rstrip("/") for link in links if link]
    return " | ".join(str(b).strip() for b in bits if b and str(b).strip())


def build_cover_letter(identity: dict, cover: dict, font_name: str, docx_path: Path, txt_path: Path) -> None:
    name = str(identity.get("name") or "").strip()
    today = dt.date.today()
    date_text = f"{today:%B} {today.day}, {today.year}"
    greeting = str(cover.get("greeting") or "Dear Hiring Team,").strip()
    closing_lines = [ln.strip() for ln in str(cover.get("closing") or "Best regards,").splitlines() if ln.strip()]
    if name and (not closing_lines or closing_lines[-1] != name):
        closing_lines.append(name)
    paragraphs = [" ".join(str(p).split()) for p in cover["paragraphs"] if str(p).strip()]

    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    head = d.add_paragraph()
    run = head.add_run(name)
    run.bold = True
    run.font.size = Pt(16)
    head.paragraph_format.space_after = Pt(2)
    d.add_paragraph(contact_line(identity)).paragraph_format.space_after = Pt(14)
    d.add_paragraph(date_text)
    d.add_paragraph(greeting)
    for p in paragraphs:
        d.add_paragraph(p)
    closing = d.add_paragraph()
    for i, line in enumerate(closing_lines):
        if i:
            closing.add_run().add_break()
        closing.add_run(line)
    d.save(str(docx_path))

    text = "\n\n".join([name + "\n" + contact_line(identity), date_text, greeting, *paragraphs, "\n".join(closing_lines)])
    txt_path.write_text(text + "\n", encoding="utf-8")


# ----------------------------------------------------------------- PDF

def find_soffice() -> str | None:
    env = os.environ.get("SOFFICE_PATH")
    if env and Path(env).exists():
        return env
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for cand in (
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/usr/lib/libreoffice/program/soffice",
        "/opt/libreoffice/program/soffice",
        "/snap/bin/libreoffice",
    ):
        if Path(cand).exists():
            return cand
    return None


def to_pdf(soffice: str, src: Path, timeout: int) -> Path:
    pdf = src.with_suffix(".pdf")
    if pdf.exists():
        pdf.unlink()
    # A throwaway profile avoids clashes with a LibreOffice window the user already has open.
    with tempfile.TemporaryDirectory(prefix="lo_profile_") as profile_dir:
        cmd = [soffice, f"-env:UserInstallation={Path(profile_dir).as_uri()}", "--headless", "--norestore",
               "--nolockcheck", "--convert-to", "pdf", "--outdir", str(src.parent), str(src)]
        try:
            res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            raise RuntimeError(f"LibreOffice timed out after {timeout} s converting {src.name}") from None
    if res.returncode != 0 or not pdf.exists():
        detail = (res.stderr or res.stdout or "").strip()[-400:]
        raise RuntimeError(f"LibreOffice could not convert {src.name} (exit {res.returncode}). {detail}")
    return pdf


def page_count(pdf: Path) -> int:
    return len(PdfReader(str(pdf)).pages)


# ----------------------------------------------------------------- dump

def dump(path: Path, as_json: bool) -> int:
    if not path.is_file():
        fail(EXIT_USAGE, f"file not found: {path}")
    d = docx.Document(str(path))
    rows = []
    for i, para in enumerate(all_paragraphs(d)):
        text = normalize(para.text)
        if text:
            rows.append({"index": i, "location": para.location, "style": para.style, "text": text})
    if as_json:
        print(json.dumps(rows, indent=2, ensure_ascii=False))
    else:
        for r in rows:
            style = f" ({r['style']})" if r["style"] else ""
            print(f"[{r['index']}] {r['location']}{style}: {r['text']}")
    return EXIT_OK


# ----------------------------------------------------------------- main

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description="Tailor the master resume for one posting.",
        epilog="Exit codes: 0 ok; 1 usage, spec or profile error; 2 anchor not found or ambiguous (nothing written); "
               "3 resume PDF longer than resume.max_pages; 4 PDF conversion failed; 5 unsupported numbers with --strict-numbers.",
    )
    ap.add_argument("--profile", default="profile/profile.yaml", help="path to profile.yaml")
    ap.add_argument("--spec", help="path to applications/<date>/<id>/spec.json")
    ap.add_argument("--no-pdf", action="store_true", help="skip LibreOffice conversion and the page check")
    ap.add_argument("--strict-numbers", action="store_true", help="exit 5 if a number is not backed by the master resume or profile")
    ap.add_argument("--timeout", type=int, default=180, help="seconds allowed per LibreOffice conversion")
    ap.add_argument("--dump", nargs="?", const="", metavar="DOCX",
                    help="list every paragraph (body, tables, text boxes, headers, footers) with normalized text and exit; "
                         "defaults to resume.master_path from --profile")
    ap.add_argument("--json", action="store_true", help="with --dump: JSON output")
    args = ap.parse_args(argv)

    profile_path = Path(os.path.expanduser(args.profile))
    if args.dump is not None:
        if args.dump:
            return dump(Path(os.path.expanduser(args.dump)), args.json)
        profile = load_yaml(profile_path)
        return dump(resolve_master(profile_path, (profile.get("resume") or {}).get("master_path") or ""), args.json)
    if not args.spec:
        ap.error("--spec is required (or use --dump)")

    profile = load_yaml(profile_path)
    identity = profile.get("identity") or {}
    resume_cfg = profile.get("resume") or {}
    anchors = resume_cfg.get("anchors") or {}
    name = str(identity.get("name") or "").strip()
    if not name:
        fail(EXIT_USAGE, "identity.name is empty in profile.yaml")
    if not resume_cfg.get("master_path"):
        fail(EXIT_USAGE, "resume.master_path is empty in profile.yaml")
    try:
        max_pages = int(resume_cfg.get("max_pages") or 2)
    except (TypeError, ValueError):
        fail(EXIT_USAGE, "resume.max_pages must be a whole number")
    master_path = resolve_master(profile_path, resume_cfg["master_path"])

    spec_path = Path(os.path.expanduser(args.spec))
    spec = load_spec(spec_path)
    company, role = str(spec["company"]).strip(), str(spec["role"]).strip()
    out_dir = Path(os.path.expanduser(spec["out_dir"])) if spec.get("out_dir") else spec_path.parent
    summary: dict = {
        "ok": False, "company": company, "role": role, "job_url": spec.get("job_url", ""),
        "master": str(master_path), "out_dir": str(out_dir), "files": {}, "replacements": [],
        "missing_anchors": [], "numbers_guard": {}, "pages": None, "max_pages": max_pages,
        "page_check": "skipped", "warnings": [],
    }

    document = docx.Document(str(master_path))
    paras = all_paragraphs(document)
    master_text = "\n".join(p.text for p in paras)
    try:
        font_name = document.styles["Normal"].font.name or "Calibri"
    except KeyError:
        font_name = "Calibri"

    # 1. Locate and replace every anchor; write nothing if any anchor is missing or ambiguous.
    for rep in spec["replacements"]:
        key, new_text = rep["anchor"], rep["text"]
        anchor_text = str(anchors.get(key) or "").strip()
        if not anchor_text:
            summary["missing_anchors"].append({"anchor": key, "reason": f"resume.anchors.{key} is not set in profile.yaml"})
            continue
        kind, matches = find_matches(paras, anchor_text)
        if kind not in ("exact", "substring"):
            reason = {
                "none": "text not found in the master resume",
                "fragment": (f"found only as a fragment inside a longer paragraph; an anchor must be the whole paragraph, "
                             f"or at least {MIN_SUBSTRING} characters at its start or end (e.g. the text after a 'Skills:' label)"),
                "ambiguous": (f"anchor text occurs in {len([m for m in matches if m[0].primary])} places; "
                              "make it longer or unique"),
            }[kind]
            summary["missing_anchors"].append({"anchor": key, "reason": reason, "anchor_text": anchor_text})
            continue
        for para, start, end in matches:
            replace_range(para, start, end, new_text)
            if normalize(new_text) not in normalize(para.text):
                fail(EXIT_USAGE, f"internal check failed after replacing '{key}'", summary)
        primary_hits = [m for m in matches if m[0].primary]
        summary["replacements"].append({"anchor": key, "match": kind, "count": len(matches),
                                        "locations": [m[0].location for m in matches]})
        if len(primary_hits) > 1:
            summary["warnings"].append(f"anchor '{key}' matched {len(primary_hits)} paragraphs; all were replaced")
        if key == "headline" and len(new_text) > HEADLINE_MAX:
            summary["warnings"].append(f"headline is {len(new_text)} characters (keep it at or under {HEADLINE_MAX})")
        if len(new_text) > 1.25 * len(anchor_text) + 10:
            summary["warnings"].append(f"'{key}' is {len(new_text)} characters vs {len(anchor_text)} in the master; watch the page count")

    if summary["missing_anchors"]:
        for m in summary["missing_anchors"]:
            log(f"ANCHOR NOT FOUND: {m['anchor']} - {m['reason']}")
        log("No files were written. Fix resume.anchors in profile.yaml (copy the text from --dump) or the spec, then re-run.")
        summary["exit_code"] = EXIT_ANCHOR
        print(json.dumps(summary, indent=2))
        return EXIT_ANCHOR

    # 2. Numbers guard: every figure must already exist in the master resume or profile.yaml.
    profile_text = profile_path.read_text(encoding="utf-8")
    guard = numbers_guard(spec, master_text + "\n" + profile_text)
    summary["numbers_guard"] = guard
    for item in guard["unsupported"]:
        log(f"WARNING numbers guard: {item['number']} in {item['where']} is not in the master resume or profile.yaml")

    # 3. Save the tailored resume and the optional cover letter.
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = safe_name(spec.get("file_stem") or f"{name} - Resume - {company}")
    resume_docx = out_dir / f"{stem}.docx"
    document.save(str(resume_docx))
    summary["files"]["resume_docx"] = str(resume_docx)
    cover_docx = None
    if spec.get("cover_letter"):
        cover_stem = safe_name(f"{name} - Cover Letter - {company}")
        cover_docx, cover_txt = out_dir / f"{cover_stem}.docx", out_dir / f"{cover_stem}.txt"
        build_cover_letter(identity, spec["cover_letter"], font_name, cover_docx, cover_txt)
        summary["files"].update(cover_docx=str(cover_docx), cover_txt=str(cover_txt))
        words = sum(len(str(p).split()) for p in spec["cover_letter"]["paragraphs"])
        if words > 250:
            summary["warnings"].append(f"cover letter body is {words} words (aim for about 150)")

    # 4. PDF conversion and page check.
    code = EXIT_OK
    if not args.no_pdf:
        soffice = find_soffice()
        if not soffice:
            summary["exit_code"] = EXIT_PDF
            summary["error"] = ("LibreOffice (soffice) not found. Install LibreOffice, set SOFFICE_PATH, "
                                "or re-run with --no-pdf to skip the PDF and page check.")
            log("ERROR: " + summary["error"])
            print(json.dumps(summary, indent=2))
            return EXIT_PDF
        try:
            resume_pdf = to_pdf(soffice, resume_docx, args.timeout)
            summary["files"]["resume_pdf"] = str(resume_pdf)
            if cover_docx:
                cover_pdf = to_pdf(soffice, cover_docx, args.timeout)
                summary["files"]["cover_pdf"] = str(cover_pdf)
                if page_count(cover_pdf) > 1:
                    summary["warnings"].append("cover letter PDF is longer than one page")
        except RuntimeError as exc:
            summary["exit_code"] = EXIT_PDF
            summary["error"] = str(exc)
            log(f"ERROR: {exc}")
            print(json.dumps(summary, indent=2))
            return EXIT_PDF
        pages = page_count(resume_pdf)
        summary["pages"] = pages
        if pages > max_pages:
            summary["page_check"] = "fail"
            log(f"ERROR: the tailored resume PDF has {pages} pages; the limit is {max_pages} (resume.max_pages). "
                "Shorten the replacement texts (headline, summary bullets, skills line) and re-run.")
            code = EXIT_PAGES
        else:
            summary["page_check"] = "pass"

    if code == EXIT_OK and args.strict_numbers and guard["unsupported"]:
        log("ERROR: numbers guard failed (--strict-numbers). Use figures exactly as the master resume states them, or drop them.")
        code = EXIT_NUMBERS
    for w in summary["warnings"]:
        log(f"WARNING: {w}")
    summary["ok"] = code == EXIT_OK
    summary["exit_code"] = code
    print(json.dumps(summary, indent=2))
    return code


if __name__ == "__main__":
    sys.exit(main())
